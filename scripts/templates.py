"""Generate editable design inputs; never overwrite an existing project file."""
from __future__ import annotations
import io
from pathlib import Path
from common import SCHEMAS, encode, transaction, read_json

ASSETS = Path(__file__).resolve().parent.parent / "assets"

def workbook_bytes():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    for kind, (title, headers) in SCHEMAS.items():
        ws = wb.create_sheet(title)
        ws.append(headers)
        ws.freeze_panes = "B2"
        ws.sheet_view.showGridLines = False
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}31"
        for col, header in enumerate(headers, 1):
            letter = get_column_letter(col)
            ws.column_dimensions[letter].width = 46 if "描述" in header else 32 if header in ("备注", "关联卡牌") else 22 if header in ("key", "名称", "关键词") else 16
            cell = ws.cell(1, col)
            cell.font = Font(name="Microsoft YaHei", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="334155")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            for row in range(2, 32):
                cell = ws.cell(row, col)
                cell.font = Font(name="Microsoft YaHei", size=11, color="1F2937")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.fill = PatternFill("solid", fgColor="FFF9E8" if row % 2 == 0 else "FFFFFF")
                if header in ("费用", "升级后费用"):
                    cell.number_format = "0"
        ws.row_dimensions[1].height = 28
        for row in range(2, 32):
            ws.row_dimensions[row].height = 44
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

def content_files(root, workbook):
    return {
        workbook: workbook_bytes(),
        root / "design/mechanics.md": (ASSETS / "mechanics.md").read_bytes(),
        root / "design/character.md": (ASSETS / "character-design.md").read_bytes(),
        root / "design/filling-guide.md": (ASSETS / "filling-guide.md").read_bytes(),
    }

def init_project(root, profile="common"):
    if profile not in ("common", "character"):
        raise ValueError("Expected common or character profile")
    root = Path(root).resolve()
    config = read_json(ASSETS / "project-config.json")
    if profile == "character":
        config["required_checks"] += ["tutorials", "sts2_agent", "mcp_server", "openpyxl", "Pillow", "official_reference", "project", "ritsulib"]
    files = {
        root / "moddev.json": encode(config),
        root / "design/project.md": (ASSETS / "project-design.md").read_bytes(),
        root / ".moddev/progress.json": encode({"schema_version": 1, "preferences": {}, "entries": {}}),
    }
    if profile == "character":
        files.update(content_files(root, root / config["workbook"]))
        files[root / config["asset_manifest"]] = encode({"schema_version": 1, "entries": []})
    transaction(files, create_only=True)
    return {"profile": profile, "created": [str(p.relative_to(root)) for p in files]}

def add_content(project):
    """Fill missing content inputs while preserving existing user-authored files."""
    candidates = [
        (project.workbook, workbook_bytes),
        (project.root / "design/mechanics.md", lambda: (ASSETS / "mechanics.md").read_bytes()),
        (project.root / "design/character.md", lambda: (ASSETS / "character-design.md").read_bytes()),
        (project.root / "design/filling-guide.md", lambda: (ASSETS / "filling-guide.md").read_bytes()),
        (project.manifest, lambda: encode({"schema_version": 1, "entries": []})),
    ]
    paths = [path.resolve() for path, _ in candidates]
    if len(set(paths)) != len(paths):
        raise ValueError("Content template destinations must be distinct")
    files, preserved = {}, []
    for path, generate in candidates:
        if path.exists():
            if not path.is_file():
                raise ValueError(f"Expected a file destination, found directory: {path}")
            preserved.append(str(path.relative_to(project.root)))
        else:
            files[path] = generate()
    transaction(files, create_only=True)
    return {"created": [str(p.relative_to(project.root)) for p in files],
            "preserved": preserved}
