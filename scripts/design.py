"""Read-only workbook comparison and explicit evidence-backed baseline acceptance."""
from __future__ import annotations
import re
from openpyxl import load_workbook
from common import SCHEMAS, digest, encode, file_hash, read_json, stamp, transaction

KEY = re.compile(r"^[A-Z][A-Za-z0-9]*$")
REQUIRED = {
    "cards": ["key", "名称", "稀有度", "类型", "费用", "描述"],
    "keywords": ["关键词", "描述"],
    "powers": ["key", "名称", "描述"],
    "relics": ["key", "名称", "遗物池", "稀有度", "描述"],
}

def read_design(project):
    wb = load_workbook(project.workbook, read_only=True, data_only=False)
    result, issues = {}, []
    try:
        for kind, (sheet, headers) in SCHEMAS.items():
            result[kind] = {}
            if sheet not in wb.sheetnames:
                issues.append({"kind": kind, "error": f"Missing sheet: {sheet}"})
                continue
            rows = wb[sheet].iter_rows()
            actual = [c.value for c in next(rows, [])]
            while actual and actual[-1] is None:
                actual.pop()
            if actual != headers:
                issues.append({"kind": kind, "error": f"Expected columns: {headers}; found: {actual}"})
                continue
            for line, cells in enumerate(rows, 2):
                values = [c.value for c in cells[:len(headers)]]
                if not any(v is not None and v != "" for v in values):
                    continue
                values += [None] * (len(headers) - len(values))
                row = dict(zip(headers, values))
                identity = row[headers[0]]
                identity = str(identity).strip() if identity is not None else ""
                errors = []
                if any(c.data_type == "f" for c in cells):
                    errors.append("Design cells must contain literal values, not formulas")
                for field in REQUIRED[kind]:
                    if row[field] is None or (isinstance(row[field], str) and not row[field].strip()):
                        errors.append(f"Missing field: {field}")
                if not identity or (kind != "keywords" and not KEY.fullmatch(identity)):
                    errors.append("Expected a PascalCase key" if kind != "keywords" else "Missing keyword")
                if identity in result[kind]:
                    errors.append(f"Duplicate identifier: {identity}")
                row[headers[0]] = identity
                if kind == "relics" and row["遗物池"] not in project.data.get("pools", {}):
                    errors.append(f"Unknown relic pool: {row['遗物池']}")
                if kind == "cards":
                    for field in ("费用", "升级后费用"):
                        v = row[field]
                        if v is not None and (isinstance(v, bool) or not isinstance(v, (int, float)) and v != "X"):
                            errors.append(f"{field}: use a number or X")
                result[kind][identity] = row
                issues.extend({"kind": kind, "key": identity, "row": line, "error": error} for error in errors)
    finally:
        wb.close()
    return result, issues

def linked_card(value, cards):
    if value is None or str(value).strip() in ("", "无"):
        return None
    value = str(value).strip()
    match = re.search(r"[（(]([A-Z][A-Za-z0-9]*)[）)]$", value)
    if match:
        key = match.group(1)
        if key in cards:
            return key
        raise ValueError(f"Unknown linked card key: {key}")
    if value in cards:
        return value
    hits = [key for key, row in cards.items() if row.get("名称") == value]
    if len(hits) != 1:
        raise ValueError(f"Linked card must match exactly one key/name: {value}; candidates={hits}")
    return hits[0]

def differences(project, kinds=None, keys=None):
    data, issues = read_design(project)
    kinds = kinds or list(SCHEMAS)
    reports = {}
    for kind in kinds:
        current = data[kind]
        old = project.baseline(kind)["entries"]
        changes = []
        for key in sorted(set(current) | set(old)):
            if keys and key not in keys:
                continue
            before, after = old.get(key), current.get(key)
            status = "added" if before is None else "deleted" if after is None else "modified" if before != after else "unchanged"
            fields = {field: {"before": (before or {}).get(field), "after": (after or {}).get(field)}
                      for field in SCHEMAS[kind][1] if (before or {}).get(field) != (after or {}).get(field)}
            if status != "unchanged":
                changes.append({"key": key, "status": status, "fields": fields})
        added = [x["key"] for x in changes if x["status"] == "added"]
        removed = [x["key"] for x in changes if x["status"] == "deleted"]
        name_field = "关键词" if kind == "keywords" else "名称"
        rename = [{"from": a, "to": b} for a in removed for b in added
                  if old[a].get(name_field) == current[b].get(name_field)
                  or (old[a].get("描述") and old[a].get("描述") == current[b].get("描述"))]
        reports[kind] = {"changes": changes, "possible_renames": rename,
                         "baseline_exists": project.baseline_path(kind).exists()}
    for kind in kinds:
        for key, row in data[kind].items():
            if keys and key not in keys:
                continue
            if kind == "powers":
                try:
                    linked = linked_card(row.get("关联卡牌"), data["cards"])
                    if linked and any(i["kind"] == "cards" and i.get("key") in (None, linked) for i in issues):
                        raise ValueError(f"Linked card design is invalid: {linked}")
                except ValueError as error:
                    issues.append({"kind": kind, "key": key, "error": str(error)})
    selected_issues = [i for i in issues if i["kind"] in kinds and (not keys or "key" not in i or i["key"] in keys)]
    for kind in kinds:
        for key in keys or []:
            if key not in data[kind] and key not in project.baseline(kind)["entries"]:
                selected_issues.append({"kind": kind, "key": key, "error": "Unknown identifier"})
    return {"reports": reports, "issues": selected_issues}, data

def entry(project, kind, key):
    report, data = differences(project, [kind], [key])
    if report["issues"]:
        raise ValueError(str(report["issues"]))
    if key not in data[kind]:
        raise ValueError("Deletion is report-only; use an explicit compatibility migration")
    return data[kind][key]

def hash_paths(project, paths):
    if not paths:
        raise ValueError("At least one implementation file and one check report are required")
    result = {}
    for value in paths:
        path = project.path(value)
        if not path.is_file() or path.stat().st_size == 0:
            raise ValueError(f"Missing or empty evidence file: {path}")
        result[str(path.relative_to(project.root))] = file_hash(path)
    return result

def record(project, kind, key, artifacts, checks, behavior, reason, implementation_id=None):
    row = entry(project, kind, key)
    if behavior in ("reused", "skipped") and not reason.strip():
        raise ValueError("Reused/skipped behavior requires a reason and evidence in the check report")
    if kind == "keywords" and not KEY.fullmatch(implementation_id or ""):
        raise ValueError("Keyword verification requires --implementation-id PascalCase")
    if kind == "keywords":
        bindings = project.data.get("keyword_ids", {})
        if bindings.get(key) != implementation_id:
            raise ValueError("Set keyword_ids in project config before recording verification")
        if any(name != key and value == implementation_id for name, value in bindings.items()):
            raise ValueError("Duplicate keyword implementation ID; resolve explicit rename mapping first")
    progress = read_json(project.progress, {"schema_version": 1, "preferences": {}, "entries": {}})
    old = progress["entries"].get(kind, {}).get(key, {})
    if old.get("implementation_id") and old["implementation_id"] != implementation_id:
        raise ValueError("Stable implementation ID changed; handle explicit compatibility migration first")
    evidence = {
        "status": "verified", "design_hash": digest(row), "context_hash": project.context_hash(),
        "artifacts": hash_paths(project, artifacts), "checks": hash_paths(project, checks),
        "behavior": behavior, "reason": reason, "implementation_id": implementation_id,
        "recorded_at": stamp(),
    }
    if set(evidence["artifacts"]) & set(evidence["checks"]):
        raise ValueError("Check reports must be separate from implementation files")
    progress["entries"].setdefault(kind, {})[key] = evidence
    transaction({project.progress: encode(progress)})
    return evidence

def accept(project, kind, keys=None):
    report, data = differences(project, [kind], keys)
    if report["issues"]:
        raise ValueError(str(report["issues"]))
    changes = report["reports"][kind]["changes"]
    if any(c["status"] == "deleted" for c in changes):
        raise ValueError("Deletion is report-only. Select completed remaining entries with --key.")
    selected = keys or [c["key"] for c in changes]
    if not selected:
        return {"accepted": [], "kind": kind}
    progress = read_json(project.progress, {"entries": {}})
    current = project.baseline(kind)
    accepted = dict(current["entries"])
    evidence_copy = {}
    for key in selected:
        row = entry(project, kind, key)
        evidence = progress["entries"].get(kind, {}).get(key, {})
        if evidence.get("status") != "verified" or evidence.get("design_hash") != digest(row):
            raise ValueError(f"{kind}/{key}: implementation/verification is missing or stale")
        if evidence.get("context_hash") != project.context_hash():
            raise ValueError(f"{kind}/{key}: project configuration changed; review and record again")
        if evidence.get("behavior") not in ("passed", "reused", "skipped"):
            raise ValueError(f"{key}: missing behavior verification disposition")
        if evidence["behavior"] in ("reused", "skipped") and not evidence.get("reason", "").strip():
            raise ValueError(f"{key}: missing reuse/skip reason")
        if kind == "keywords" and project.data.get("keyword_ids", {}).get(key) != evidence.get("implementation_id"):
            raise ValueError(f"{key}: keyword implementation binding is missing")
        for section in ("artifacts", "checks"):
            if not evidence.get(section):
                raise ValueError(f"{key}: missing {section}")
            for path, expected in evidence[section].items():
                if not project.path(path).is_file() or file_hash(project.path(path)) != expected:
                    raise ValueError(f"{key}: evidence file changed: {path}")
        accepted[key] = row
        evidence_copy[key] = evidence
    result = {"schema_version": 1, "accepted_at": stamp(), "entries": accepted}
    history = project.baseline_path(kind).parent / "history" / (stamp() + ".json")
    transaction({history: encode({"previous": current, "accepted": selected, "evidence": evidence_copy}),
                 project.baseline_path(kind): encode(result)})
    return {"accepted": selected, "kind": kind}

