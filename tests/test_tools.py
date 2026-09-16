"""Isolated unit and CLI integration checks; never touches a game or user workbook."""
from __future__ import annotations
import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from openpyxl import load_workbook
from PIL import Image
from common import Project, encode, file_hash, read_json, transaction
from templates import init_project
from design import read_design, differences, linked_card, record, accept
from art import sync, icons, placeholder
from doctor import doctor

class ToolTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="sts2-skill-test-")
        self.root = Path(self.tmp.name)
        init_project(self.root, profile="character")
        self.config = self.root / "moddev.json"
        self.configure(keyword_ids={"余势": "MomentumTip"})
        self.project = Project(self.config)

    def tearDown(self):
        self.tmp.cleanup()

    def configure(self, **kwargs):
        config = read_json(self.config)
        config.update(kwargs)
        self.config.write_bytes(encode(config))
        self.project = Project(self.config)

    def rows(self, **entries):
        wb = load_workbook(self.project.workbook)
        titles = {"cards": "卡牌", "keywords": "关键词", "powers": "状态", "relics": "遗物"}
        for kind, rows in entries.items():
            ws = wb[titles[kind]]
            ws.delete_rows(2, ws.max_row)
            for row in rows:
                ws.append(row)
        wb.save(self.project.workbook)
        wb.close()

    def full(self):
        self.rows(cards=[["Spark", "火花", "初始", "攻击", 1, "造成6点伤害。", 1, "造成9点伤害。", None]],
                  keywords=[["余势", "下一张攻击牌额外造成2点伤害。", None]],
                  powers=[["Spark", "余势", "伤害增加{}点。", "Spark"]],
                  relics=[["Spark", "徽章", "通用", "初始", "获得3点格挡。", None]])

    def evidence(self, kind, key):
        artifact = self.root / "Mod" / kind / "content.cs"
        report = self.root / "checks" / (kind + ".json")
        artifact.parent.mkdir(parents=True, exist_ok=True)
        report.parent.mkdir(parents=True, exist_ok=True)
        artifact.write_text("// isolated test implementation", encoding="utf-8")
        report.write_bytes(encode({"test_fixture": True, "exit_code": 0, "assertions": ["test-only evidence"]}))
        return record(self.project, kind, key, [str(artifact.relative_to(self.root))],
                      [str(report.relative_to(self.root))], "passed", "",
                      "MomentumTip" if kind == "keywords" else None)

    def image(self, relative, color="red", size=(400, 200)):
        path = self.root / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        Image.new("RGBA", size, color).save(path)
        return path

    def manifest(self, items):
        self.project.manifest.write_bytes(encode({"schema_version": 1, "entries": items}))

    def art_entry(self, kind="cards", key="Spark", target="Mod/images/Spark.png"):
        impl = self.root / "Mod" / kind / (key + ".cs")
        impl.parent.mkdir(parents=True, exist_ok=True)
        impl.write_text("// fixture", encoding="utf-8")
        return {"kind": kind, "key": key, "implementation": str(impl.relative_to(self.root)), "target": target}

    def test_template_four_empty_sheets_and_preserved_existing(self):
        wb = load_workbook(self.project.workbook)
        self.assertEqual(wb.sheetnames, ["卡牌", "关键词", "状态", "遗物"])
        self.assertEqual([wb[s].max_column for s in wb.sheetnames], [9, 3, 4, 6])
        self.assertTrue(all(wb[s].freeze_panes == "B2" for s in wb.sheetnames))
        self.assertTrue(all(wb[s]["A2"].value is None for s in wb.sheetnames))
        wb.close()
        before = file_hash(self.project.workbook)
        with self.assertRaises(FileExistsError):
            init_project(self.root, profile="character")
        self.assertEqual(before, file_hash(self.project.workbook))

    def test_all_kinds_independent_same_key_and_workbook_read_only(self):
        self.full()
        before = file_hash(self.project.workbook)
        report, data = differences(self.project)
        self.assertFalse(report["issues"])
        self.assertTrue(all(len(r["changes"]) == 1 for r in report["reports"].values()))
        for kind, key in [("cards", "Spark"), ("keywords", "余势"), ("powers", "Spark"), ("relics", "Spark")]:
            self.evidence(kind, key)
            self.assertEqual(accept(self.project, kind, [key])["accepted"], [key])
        self.assertTrue(all(not r["changes"] for r in differences(self.project)[0]["reports"].values()))
        self.assertEqual(before, file_hash(self.project.workbook))

    def test_duplicate_identifier_rejected(self):
        self.rows(powers=[["Spark", "状态", "说明", None], ["Spark", "重复", "说明", None]])
        self.assertIn("Duplicate", str(read_design(self.project)[1]))
        with self.assertRaises(ValueError):
            accept(self.project, "powers", ["Spark"])

    def test_missing_field_and_formula_reported(self):
        self.rows(cards=[["Spark", "火花", "初始", "攻击", None, "=1+2"]])
        errors = str(read_design(self.project)[1])
        self.assertIn("Missing field", errors)
        self.assertIn("literal", errors)

    def test_unknown_pool_rejected(self):
        self.rows(relics=[["Badge", "徽章", "未知角色池", "初始", "说明", None]])
        self.assertIn("Unknown relic pool", str(read_design(self.project)[1]))

    def test_notes_and_numeric_change_are_diffed(self):
        self.full()
        self.evidence("cards", "Spark")
        accept(self.project, "cards", ["Spark"])
        self.rows(cards=[["Spark", "火花", "初始", "攻击", 0, "造成6点伤害。", 1, "造成9点伤害。", "目标：自身"]])
        fields = differences(self.project, ["cards"])[0]["reports"]["cards"]["changes"][0]["fields"]
        self.assertEqual(set(fields), {"费用", "备注"})

    def test_no_evidence_no_accept(self):
        self.full()
        with self.assertRaises(ValueError):
            accept(self.project, "cards", ["Spark"])
        self.assertFalse(self.project.baseline_path("cards").exists())

    def test_changed_artifact_invalidates_evidence(self):
        self.full()
        self.evidence("cards", "Spark")
        (self.root / "Mod/cards/content.cs").write_text("changed")
        with self.assertRaisesRegex(ValueError, "changed"):
            accept(self.project, "cards", ["Spark"])

    def test_changed_report_and_config_invalidate_evidence(self):
        self.full()
        self.evidence("cards", "Spark")
        (self.root / "checks/cards.json").write_text("{}")
        with self.assertRaisesRegex(ValueError, "changed"):
            accept(self.project, "cards", ["Spark"])
        self.evidence("cards", "Spark")
        self.configure(languages=["eng"])
        with self.assertRaisesRegex(ValueError, "configuration changed"):
            accept(self.project, "cards", ["Spark"])

    def test_design_change_invalidates_record(self):
        self.full()
        self.evidence("powers", "Spark")
        self.rows(powers=[["Spark", "余势", "改成另一效果。", "Spark"]])
        with self.assertRaisesRegex(ValueError, "stale"):
            accept(self.project, "powers", ["Spark"])

    def test_keyword_rename_is_not_automatic(self):
        self.full()
        self.evidence("keywords", "余势")
        accept(self.project, "keywords", ["余势"])
        self.rows(keywords=[["蓄势", "下一张攻击牌额外造成2点伤害。", None]])
        report, _ = differences(self.project, ["keywords"])
        self.assertEqual(report["reports"]["keywords"]["possible_renames"], [{"from": "余势", "to": "蓄势"}])
        with self.assertRaises(ValueError):
            accept(self.project, "keywords")
        self.assertIn("余势", self.project.baseline("keywords")["entries"])

    def test_keyword_binding_required(self):
        self.full()
        self.configure(keyword_ids={})
        with self.assertRaisesRegex(ValueError, "keyword_ids"):
            self.evidence("keywords", "余势")

    def test_deleted_item_retained_and_other_key_can_accept(self):
        self.rows(powers=[["First", "一", "说明", None], ["Second", "二", "说明", None]])
        # Both share a stable fixture file; recording after writing all content is allowed.
        for key in ["First", "Second"]:
            self.evidence("powers", key)
        accept(self.project, "powers")
        self.rows(powers=[["Second", "二", "新说明", None]])
        self.evidence("powers", "Second")
        with self.assertRaises(ValueError):
            accept(self.project, "powers")
        accept(self.project, "powers", ["Second"])
        self.assertIn("First", self.project.baseline("powers")["entries"])

    def test_link_resolution_empty_missing_ambiguous_and_brackets(self):
        cards = {"First": {"名称": "同名"}, "Second": {"名称": "同名"}}
        self.assertIsNone(linked_card("", cards))
        self.assertEqual(linked_card("旧名（First）", cards), "First")
        for text in ["同名", "Missing"]:
            with self.assertRaises(ValueError):
                linked_card(text, cards)

    def test_missing_link_reported(self):
        self.rows(powers=[["Spark", "余势", "说明", "Missing"]])
        self.assertTrue(differences(self.project, ["powers"])[0]["issues"])

    def test_placeholder_preserves_formal_art_and_tracks_generated(self):
        formal = self.image("art/cards/Spark.png", "green")
        old = formal.read_bytes()
        for refresh in [False, True]:
            with self.assertRaises(ValueError):
                placeholder(self.project, "Spark", refresh=refresh)
        self.assertEqual(old, formal.read_bytes())
        placeholder(self.project, "NewCard", "临时卡牌")
        placeholder(self.project, "NewCard", "新临时名称", refresh=True)
        with Image.open(self.root / "art/cards/NewCard.png") as image:
            self.assertEqual(image.size, (250, 190))

    def test_sync_preview_content_update_and_backup(self):
        self.full()
        source = self.image("art/cards/Spark.png")
        item = self.art_entry()
        self.manifest([item])
        self.assertEqual(sync(self.project)["entries"][0]["status"], "would_update")
        target = self.root / item["target"]
        self.assertFalse(target.exists())
        sync(self.project, apply=True)
        self.assertEqual(source.read_bytes(), target.read_bytes())
        self.assertEqual(sync(self.project)["entries"][0]["status"], "unchanged")
        self.image("art/cards/Spark.png", "blue")
        sync(self.project, apply=True)
        self.assertEqual(source.read_bytes(), target.read_bytes())
        self.assertTrue(list((self.project.state / "backups").rglob("manifest.json")))

    def test_sync_ambiguity_keeps_target(self):
        self.full()
        self.image("art/cards/a/Spark.png")
        self.image("art/cards/b/Spark.png", "blue")
        target = self.image("Mod/images/Spark.png", "green")
        before = target.read_bytes()
        self.manifest([self.art_entry()])
        self.assertTrue(sync(self.project, apply=True)["issues"])
        self.assertEqual(before, target.read_bytes())

    def test_duplicate_target_blocks_both_entries(self):
        self.full()
        self.image("art/cards/Spark.png")
        item = self.art_entry()
        self.manifest([item, dict(item)])
        self.assertTrue(sync(self.project, apply=True)["issues"])
        self.assertFalse((self.root / item["target"]).exists())

    def test_sync_bad_png_and_raw_target_blocked(self):
        self.full()
        source = self.root / "art/cards/Spark.png"
        source.parent.mkdir(parents=True)
        source.write_bytes(b"not png")
        self.manifest([self.art_entry()])
        self.assertTrue(sync(self.project, apply=True)["issues"])
        self.image("art/cards/Spark.png")
        self.manifest([self.art_entry(target="art/cards/Spark.png")])
        self.assertTrue(sync(self.project, apply=True)["issues"])

    def test_icons_skip_unlinked_preserve_existing_and_generate(self):
        self.full()
        self.rows(powers=[["Spark", "余势", "说明", "Spark"], ["Unlinked", "无关联", "说明", None]])
        self.image("art/cards/Spark.png")
        preview = icons(self.project)
        self.assertEqual(preview["entries"][0]["status"], "would_generate")
        self.assertFalse((self.root / "art/powers/Spark.png").exists())
        icons(self.project, apply=True)
        path = self.root / "art/powers/Spark.png"
        with Image.open(path) as image:
            self.assertEqual((image.size, image.mode), ((256, 256), "RGBA"))
        before = path.read_bytes()
        self.image("art/cards/Spark.png", "blue")
        self.assertEqual(icons(self.project, apply=True)["entries"][0]["status"], "preserved_existing")
        self.assertEqual(before, path.read_bytes())
        icons(self.project, apply=True, overwrite=True)
        self.assertNotEqual(before, path.read_bytes())

    def test_icon_output_cannot_overwrite_card_source_directory(self):
        self.full()
        self.image("art/cards/Spark.png")
        before = (self.root / "art/cards/Spark.png").read_bytes()
        self.configure(asset_roots={"cards": "art/cards", "powers": "art/cards", "relics": "art/relics"})
        result = icons(self.project, apply=True, overwrite=True)
        self.assertTrue(result["issues"])
        self.assertEqual(before, (self.root / "art/cards/Spark.png").read_bytes())

    def test_sync_relic_and_power_kinds(self):
        self.full()
        for kind in ["powers", "relics"]:
            source = self.image(f"art/{kind}/Spark.png")
            item = self.art_entry(kind=kind, target=f"Mod/images/{kind}/Spark.png")
            self.manifest([item])
            result = sync(self.project, apply=True, kind=kind)
            self.assertFalse(result["issues"])
            self.assertEqual(source.read_bytes(), (self.root / item["target"]).read_bytes())

    def test_transaction_restores_old_and_removes_new_files(self):
        import common
        original = self.root / "existing.txt"
        created = self.root / "new.txt"
        failing = self.root / "fail.txt"
        original.write_bytes(b"original")
        real_write = common.atomic_write
        def fail_one(path, payload):
            if Path(path) == failing:
                raise OSError("simulated disk failure")
            return real_write(path, payload)
        with patch("common.atomic_write", side_effect=fail_one):
            with self.assertRaises(OSError):
                transaction({original: b"updated", created: b"new", failing: b"bad"},
                            self.root / "backups")
        self.assertEqual(original.read_bytes(), b"original")
        self.assertFalse(created.exists())
        self.assertTrue((self.root / "backups/manifest.json").exists())

    def test_paths_cannot_escape_project(self):
        with self.assertRaises(ValueError):
            self.project.path("../outside.png")

    def test_invalid_linked_card_blocks_state_acceptance(self):
        self.rows(cards=[["Spark", "火花", "初始", "攻击", 1, None]],
                  powers=[["Momentum", "余势", "说明", "Spark"]])
        result, _ = differences(self.project, ["powers"])
        self.assertIn("Linked card design is invalid", str(result["issues"]))

    def test_reuse_and_skip_require_reason(self):
        self.full()
        evidence = self.evidence("cards", "Spark")
        for behavior in ("reused", "skipped"):
            with self.assertRaisesRegex(ValueError, "reason"):
                record(self.project, "cards", "Spark", list(evidence["artifacts"]),
                       list(evidence["checks"]), behavior, "")

    def test_defer_art_does_not_generate_placeholder(self):
        self.configure(preferences={"missing_art": "defer"})
        self.full()
        differences(self.project)
        self.assertFalse((self.root / "art").exists())
        # Author explicitly selects no image; design/evidence tools do not require art generation.
        self.evidence("cards", "Spark")
        accept(self.project, "cards", ["Spark"])
        self.assertFalse((self.root / "art").exists())

    def test_partial_sync_applies_valid_and_keeps_missing(self):
        self.rows(cards=[["Good", "有图", "初始", "攻击", 1, "说明"],
                         ["Missing", "缺图", "初始", "攻击", 1, "说明"]])
        self.image("art/cards/Good.png")
        existing = self.image("Mod/images/Missing.png", "green")
        before = existing.read_bytes()
        self.manifest([self.art_entry(key="Good", target="Mod/images/Good.png"),
                       self.art_entry(key="Missing", target="Mod/images/Missing.png")])
        report = sync(self.project, apply=True)
        self.assertEqual(len(report["issues"]), 1)
        self.assertTrue((self.root / "Mod/images/Good.png").is_file())
        self.assertEqual(before, existing.read_bytes())

    def test_doctor_reports_missing_library_without_mutation(self):
        before = sorted(str(p.relative_to(self.root)) for p in self.root.rglob("*"))
        with patch("doctor.probe", return_value={"exit_code": 0, "output": "test version"}):
            result = doctor(self.project)
        self.assertIn("ritsulib", result["missing"])
        self.assertEqual(before, sorted(str(p.relative_to(self.root)) for p in self.root.rglob("*")))

    def test_formal_named_art_replaces_registered_key_placeholder(self):
        self.full()
        placeholder(self.project, "Spark", "火花", "攻击")
        old_placeholder = (self.root / "art/cards/Spark.png").read_bytes()
        final_art = self.image("art/cards/火花.png", "blue")
        item = self.art_entry()
        self.manifest([item])
        report = sync(self.project, apply=True)
        self.assertFalse(report["issues"])
        self.assertEqual((self.root / item["target"]).read_bytes(), final_art.read_bytes())
        self.assertEqual((self.root / "art/cards/Spark.png").read_bytes(), old_placeholder)
        icons(self.project, apply=True)
        with Image.open(self.root / "art/powers/Spark.png") as icon:
            self.assertEqual(icon.getpixel((128, 128)), (0, 0, 255, 255))

    def test_modified_key_image_remains_higher_priority_than_name(self):
        self.full()
        placeholder(self.project, "Spark")
        formal_key = self.image("art/cards/Spark.png", "green")
        self.image("art/cards/火花.png", "blue")
        item = self.art_entry()
        self.manifest([item])
        report = sync(self.project, apply=True)
        self.assertFalse(report["issues"])
        self.assertEqual((self.root / item["target"]).read_bytes(), formal_key.read_bytes())

    def test_placeholder_fallback_and_ambiguous_formal_art(self):
        self.full()
        placeholder(self.project, "Spark")
        item = self.art_entry()
        self.manifest([item])
        report = sync(self.project, apply=True)
        self.assertFalse(report["issues"])
        target = self.root / item["target"]
        before = target.read_bytes()
        self.image("art/cards/a/火花.png", "blue")
        self.image("art/cards/b/火花.png", "green")
        report = sync(self.project, apply=True)
        self.assertTrue(report["issues"])
        self.assertEqual(before, target.read_bytes())

    def test_unregistered_key_art_is_not_assumed_to_be_placeholder(self):
        self.full()
        key_image = self.image("art/cards/Spark.png", "gray")
        self.image("art/cards/火花.png", "blue")
        item = self.art_entry()
        self.manifest([item])
        report = sync(self.project, apply=True)
        self.assertFalse(report["issues"])
        self.assertEqual((self.root / item["target"]).read_bytes(), key_image.read_bytes())

    def test_icons_cli_key_ignores_unrelated_bad_row(self):
        self.full()
        self.rows(powers=[["Spark", "余势", "说明", "Spark"], ["Other", "未完成状态", None, None]])
        self.image("art/cards/Spark.png")
        tool = Path(__file__).resolve().parents[1] / "scripts/moddev.py"
        result = subprocess.run([sys.executable, "-X", "utf8", str(tool), "icons",
                                 "--config", str(self.config), "--key", "Spark", "--apply"],
                                capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
        self.assertFalse(json.loads(result.stdout)["issues"])
        self.assertTrue((self.root / "art/powers/Spark.png").is_file())
        self.assertTrue(icons(self.project, keys=["Other"])["issues"])
        self.assertTrue(icons(self.project)["issues"])

    def test_icons_key_keeps_sheet_and_linked_card_errors(self):
        self.full()
        self.rows(cards=[["Spark", "火花", "初始", "攻击", 1, None]])
        self.assertTrue(icons(self.project, keys=["Spark"])["issues"])
        wb = load_workbook(self.project.workbook)
        wb["状态"]["A1"] = "broken header"
        wb.save(self.project.workbook)
        wb.close()
        self.assertTrue(icons(self.project, keys=["Spark"])["issues"])

    def test_cli_all_kinds_and_art_roundtrip(self):
        tool = Path(__file__).resolve().parents[1] / "scripts/moddev.py"
        def cli(*args, code=0):
            result = subprocess.run([sys.executable, "-X", "utf8", str(tool), *args],
                                    capture_output=True, text=True, encoding="utf-8")
            self.assertEqual(result.returncode, code, result.stderr + result.stdout)
            return json.loads(result.stdout)
        args = ["--config", str(self.config)]
        cli("init", "--project", str(self.root / "another-project"))
        cli("diff", *args)
        cli("doctor", *args, code=2)
        self.full()
        diff = cli("diff", *args)
        self.assertEqual(set(diff["reports"]), {"cards", "keywords", "powers", "relics"})
        cli("placeholder", *args, "--key", "Spark", "--name", "火花", "--type", "攻击")
        cli("icons", *args, "--apply")
        self.manifest([self.art_entry()])
        cli("sync-art", *args, "--apply")
        for kind, key in [("cards", "Spark"), ("keywords", "余势"), ("powers", "Spark"), ("relics", "Spark")]:
            self.evidence(kind, key)
            record_args = ["record", *args, "--kind", kind, "--key", key,
                           "--artifact", f"Mod/{kind}/content.cs", "--check", f"checks/{kind}.json", "--behavior", "passed"]
            if kind == "keywords":
                record_args += ["--implementation-id", "MomentumTip"]
            cli(*record_args)
            cli("accept", *args, "--kind", kind, "--key", key)
        before = file_hash(self.project.workbook)
        self.assertFalse(any(x["changes"] for x in cli("diff", *args)["reports"].values()))
        self.assertEqual(before, file_hash(self.project.workbook))


class CommonFlowTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="sts2-common-test-")
        self.root = Path(self.tmp.name)
        self.tool = Path(__file__).resolve().parents[1] / "scripts/moddev.py"

    def tearDown(self):
        self.tmp.cleanup()

    def test_common_cli_runs_without_site_packages(self):
        result = subprocess.run([sys.executable, "-S", "-X", "utf8", str(self.tool),
                                 "init", "--project", str(self.root)],
                                capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(json.loads(result.stdout)["profile"], "common")
        files = {str(p.relative_to(self.root)).replace("\\", "/") for p in self.root.rglob("*") if p.is_file()}
        self.assertEqual(files, {"moddev.json", "design/project.md", ".moddev/progress.json"})
        self.assertFalse((self.root / "design/mod-design.xlsx").exists())
        result = subprocess.run([sys.executable, "-S", "-X", "utf8", str(self.tool),
                                 "doctor", "--config", str(self.root / "moddev.json")],
                                capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(result.returncode, 2, result.stderr)
        result = json.loads(result.stdout)
        self.assertIn("openpyxl", result["optional_missing"])
        self.assertNotIn("ritsulib", result["missing"])
        self.assertEqual(read_json(self.root / "moddev.json")["required_checks"], ["python", "game", "dotnet", "godot"])
        self.assertNotIn("tutorials", result["missing"])
        self.assertNotIn("sts2_agent", result["missing"])
        self.assertNotIn("mcp_server", result["missing"])
        self.assertIn("sts2_agent", result["optional_missing"])


    def test_add_content_cli_preserves_common_files(self):
        init_project(self.root)
        preserved = {p: p.read_bytes() for p in self.root.rglob("*") if p.is_file()}
        result = subprocess.run([sys.executable, "-X", "utf8", str(self.tool),
                                 "add-content", "--config", str(self.root / "moddev.json")],
                                capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(result.returncode, 0, result.stderr)
        for path, data in preserved.items():
            self.assertEqual(path.read_bytes(), data)
        wb = load_workbook(self.root / "design/mod-design.xlsx")
        self.assertEqual(wb.sheetnames, ["卡牌", "关键词", "状态", "遗物"])
        wb.close()
        self.assertTrue((self.root / "design/character.md").exists())
        self.assertTrue((self.root / ".moddev/assets.json").exists())

    def test_add_content_preserves_mechanism_and_character_documents(self):
        from templates import add_content
        init_project(self.root)
        preserved = {}
        for name in ("mechanics.md", "character.md"):
            path = self.root / "design" / name
            path.write_text("Confirmed user design: " + name, encoding="utf-8")
            preserved[path] = path.read_bytes()
        result = add_content(Project(self.root / "moddev.json"))
        for path, data in preserved.items():
            self.assertEqual(path.read_bytes(), data)
        self.assertTrue((self.root / "design/mod-design.xlsx").exists())
        self.assertEqual(len(result["preserved"]), 2)
        before = {path: path.read_bytes() for path in self.root.rglob("*") if path.is_file()}
        with patch("templates.workbook_bytes", side_effect=AssertionError("must not regenerate")):
            again = add_content(Project(self.root / "moddev.json"))
        self.assertEqual(again["created"], [])
        self.assertEqual(len(again["preserved"]), 5)
        self.assertEqual(before, {path: path.read_bytes() for path in self.root.rglob("*") if path.is_file()})

    def test_add_content_preserves_existing_asset_manifest(self):
        from templates import add_content
        init_project(self.root)
        manifest = self.root / ".moddev/assets.json"
        manifest.write_bytes(encode({"schema_version": 1, "entries": [{"custom": "keep"}]}))
        before = manifest.read_bytes()
        add_content(Project(self.root / "moddev.json"))
        self.assertEqual(before, manifest.read_bytes())

    def test_character_cli_still_generates_complete_inputs(self):
        result = subprocess.run([sys.executable, "-X", "utf8", str(self.tool),
                                 "init", "--project", str(self.root), "--profile", "character"],
                                capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(json.loads(result.stdout)["profile"], "character")
        self.assertTrue((self.root / "design/mod-design.xlsx").is_file())
        self.assertIn("ritsulib", read_json(self.root / "moddev.json")["required_checks"])

    def test_required_checks_and_legacy_compatibility(self):
        init_project(self.root)
        config_path = self.root / "moddev.json"
        data = read_json(config_path)
        data["required_checks"] = ["python"]
        config_path.write_bytes(encode(data))
        with patch("doctor.probe", return_value={"exit_code": 0, "output": "test"}):
            result = doctor(Project(config_path))
        self.assertFalse(result["missing"])
        self.assertIn("ritsulib", result["optional_missing"])
        del data["required_checks"]
        config_path.write_bytes(encode(data))
        with patch("doctor.probe", return_value={"exit_code": 0, "output": "test"}):
            result = doctor(Project(config_path))
        self.assertIn("ritsulib", result["missing"])

    def test_add_content_preserves_existing_workbook_without_regeneration(self):
        from templates import add_content
        init_project(self.root)
        workbook = self.root / "design/mod-design.xlsx"
        workbook.write_bytes(b"user workbook preserved verbatim")
        with patch("templates.workbook_bytes", side_effect=AssertionError("must not regenerate")):
            result = add_content(Project(self.root / "moddev.json"))
        self.assertEqual(workbook.read_bytes(), b"user workbook preserved verbatim")
        self.assertIn("design/mod-design.xlsx", [s.replace("\\", "/") for s in result["preserved"]])

    def test_add_content_directory_conflict_creates_no_partial_files(self):
        from templates import add_content
        init_project(self.root)
        (self.root / "design/mechanics.md").mkdir()
        with self.assertRaises(ValueError):
            add_content(Project(self.root / "moddev.json"))
        self.assertFalse((self.root / "design/mod-design.xlsx").exists())
        self.assertFalse((self.root / "design/character.md").exists())

    def test_godot_directory_missing_and_non_executable_are_rejected(self):
        from doctor import check_godot
        text_file = self.root / "godot.txt"
        text_file.write_text("4.5.1.stable.mono.official")
        with patch("doctor.probe", side_effect=AssertionError("must not execute")):
            for path in (None, self.root, self.root / "missing.exe", text_file):
                with self.subTest(path=path):
                    self.assertFalse(check_godot(path)["available"])

    def test_godot_requires_successful_mono_version(self):
        from doctor import check_godot
        candidate = self.root / "godot.exe"
        candidate.write_bytes(b"test fixture; execution mocked")
        candidate.chmod(0o755)
        cases = [
            ({"exit_code": 0, "output": "4.5.1.stable.mono.official.f62fdbde1"}, True),
            ({"exit_code": 0, "output": "4.5.1.stable.official.f62fdbde1"}, False),
            ({"exit_code": 0, "output": "not Godot; file named mono"}, False),
            ({"exit_code": 1, "output": "4.5.1.stable.mono.official.f62fdbde1"}, False),
            ({"error": "timeout"}, False),
        ]
        for output, expected in cases:
            with self.subTest(output=output), patch("doctor.probe", return_value=output):
                result = check_godot(candidate)
                self.assertEqual(result["available"], expected)
                self.assertEqual(result["dotnet_supported"], expected)

    def test_doctor_godot_directory_is_reported_missing(self):
        init_project(self.root)
        config_path = self.root / "moddev.json"
        config = read_json(config_path)
        config["environment"]["godot"] = str(self.root)
        config["required_checks"] = ["godot"]
        config_path.write_bytes(encode(config))
        with patch("doctor.probe", return_value={"exit_code": 0, "output": "test"}):
            report = doctor(Project(config_path))
        self.assertEqual(report["missing"], ["godot"])

    def test_unknown_required_check_rejected(self):
        init_project(self.root)
        config_path = self.root / "moddev.json"
        data = read_json(config_path)
        data["required_checks"] = ["misspelled-tool"]
        config_path.write_bytes(encode(data))
        with patch("doctor.probe", return_value={"exit_code": 0, "output": "test"}):
            with self.assertRaises(ValueError):
                doctor(Project(config_path))


class PlatformTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix="sts2-platform-test-")
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_posix_extensionless_godot(self):
        from doctor import check_godot
        candidate = self.root / "godot"
        candidate.write_text("fixture", encoding="utf-8")
        candidate.chmod(0o755)
        with patch("doctor.platform.system", return_value="Linux"), patch("doctor.os.access", return_value=True), patch("doctor.probe", return_value={"exit_code": 0, "output": "4.5.1.stable.mono.official.hash"}) as probe:
            self.assertTrue(check_godot(candidate)["available"])
            probe.assert_called_once_with([str(candidate), "--version"])

    def test_posix_file_requires_execute_permission(self):
        from doctor import check_godot
        candidate = self.root / "godot"
        candidate.write_text("fixture", encoding="utf-8")
        with patch("doctor.platform.system", return_value="Linux"), patch("doctor.os.access", return_value=False), patch("doctor.probe") as probe:
            self.assertFalse(check_godot(candidate)["available"])
            probe.assert_not_called()

    def test_macos_app_resolves_bundle_executable(self):
        import plistlib
        from doctor import check_godot
        bundle = self.root / "Godot Mono.app"
        binary = bundle / "Contents/MacOS/CustomGodot"
        binary.parent.mkdir(parents=True)
        binary.write_text("fixture", encoding="utf-8")
        with (bundle / "Contents/Info.plist").open("wb") as stream:
            plistlib.dump({"CFBundleExecutable": "CustomGodot"}, stream)
        with patch("doctor.platform.system", return_value="Darwin"), patch("doctor.os.access", return_value=True), patch("doctor.probe", return_value={"exit_code": 0, "output": "4.5.1.stable.mono.official.hash"}) as probe:
            result = check_godot(bundle)
            self.assertTrue(result["available"])
            self.assertEqual(result["path"], str(binary))
            probe.assert_called_once_with([str(binary), "--version"])

    def test_macos_missing_bundle_binary_is_not_a_valid_directory(self):
        from doctor import check_godot
        bundle = self.root / "Godot.app"
        bundle.mkdir()
        with patch("doctor.platform.system", return_value="Darwin"), patch("doctor.probe") as probe:
            self.assertFalse(check_godot(bundle)["available"])
            probe.assert_not_called()

    def test_configured_dotnet_works_without_path_discovery(self):
        init_project(self.root)
        config_path = self.root / "moddev.json"
        config = read_json(config_path)
        config["environment"]["dotnet"] = "tools/dotnet"
        config["required_checks"] = ["dotnet"]
        config_path.write_bytes(encode(config))
        with patch("doctor.shutil.which", return_value=None), patch("doctor.probe", return_value={"exit_code": 0, "output": "9.0.100"}) as probe:
            result = doctor(Project(config_path))
        self.assertEqual(result["missing"], [])
        self.assertEqual(probe.call_args.args[0], [str(self.root / "tools/dotnet"), "--version"])
        self.assertIn("system", result["platform"])
        self.assertIn("architecture", result["platform"])

    def test_backslash_paths_are_read_and_traversal_is_still_blocked(self):
        init_project(self.root)
        project = Project(self.root / "moddev.json")
        self.assertEqual(project.path(r"art\cards\Spark.png"), self.root / "art/cards/Spark.png")
        with self.assertRaises(ValueError):
            project.path(r"..\outside.png")

    def test_legacy_placeholder_registry_can_be_refreshed_portably(self):
        init_project(self.root)
        project = Project(self.root / "moddev.json")
        placeholder(project, "Spark")
        registry_path = project.state / "placeholders.json"
        registry = read_json(registry_path)
        self.assertIn("art/cards/Spark.png", registry)
        registry_path.write_bytes(encode({r"art\cards\Spark.png": registry["art/cards/Spark.png"]}))
        placeholder(project, "Spark", refresh=True)
        registry = read_json(registry_path)
        self.assertEqual(set(registry), {"art/cards/Spark.png"})

    def test_explicit_font_and_platform_font_fallback(self):
        from art import font, font_candidates
        for system in ("Windows", "Darwin", "Linux"):
            with self.subTest(system=system), patch("art.platform.system", return_value=system):
                self.assertTrue(font_candidates())
                with patch.object(Path, "is_file", return_value=False), patch("art.ImageFont.load_default", return_value="fallback"):
                    self.assertEqual(font(12), "fallback")
        with patch("art.ImageFont.truetype", return_value="selected") as selected:
            self.assertEqual(font(20, str(self.root / "custom.ttf")), "selected")
            selected.assert_called_once_with(str(self.root / "custom.ttf"), 20)

    @unittest.skipIf(sys.platform == "win32", "requires native POSIX process execution")
    def test_native_posix_executable_probe(self):
        from doctor import check_godot
        candidate = self.root / "godot-fixture"
        candidate.write_text("#!/bin/sh\nprintf '4.5.1.stable.mono.official.fixture\\n'\n", encoding="utf-8")
        candidate.chmod(0o755)
        self.assertTrue(check_godot(candidate)["available"])
        candidate.chmod(0o644)
        self.assertFalse(check_godot(candidate)["available"])

if __name__ == "__main__":
    unittest.main()

